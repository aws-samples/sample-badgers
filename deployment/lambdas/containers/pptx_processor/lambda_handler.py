"""Native-first PowerPoint intake and selective slide rendering."""

from __future__ import annotations

import base64
import json
import logging
import os
import re
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import boto3
import fitz
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from foundation import job_state

logger = logging.getLogger()
logger.setLevel(getattr(logging, os.environ.get("LOGGING_LEVEL", "INFO").upper()))

_SEMANTIC = re.compile(
    r"\b(chart|diagram|flow|process|workflow|architecture|map|matrix|"
    r"timeline|model|framework)\b",
    re.IGNORECASE,
)


@dataclass
class SlideFacts:
    number: int
    text_items: list[dict[str, Any]] = field(default_factory=list)
    image_count: int = 0
    image_area: int = 0
    chart_count: int = 0
    table_count: int = 0
    connector_count: int = 0
    group_count: int = 0
    smartart_count: int = 0
    warnings: list[str] = field(default_factory=list)
    notes: str = ""
    slide_area: int = 1

    @property
    def text(self) -> str:
        return "\n".join(item["text"] for item in self.text_items)


def lambda_handler(event: dict[str, Any], context: Any) -> dict[str, Any]:  # noqa: ARG001
    """Plan a PPTX deck, or render the plan-selected slides only."""
    job_id = ""
    subtask = ""
    try:
        body = json.loads(event["body"]) if "body" in event else event
        operation = str(body.get("operation") or "plan").lower()
        document_path = str(body.get("document_path") or "").strip()
        session_id = str(body.get("session_id") or "no_session")
        job_id = str(body.get("job_id") or "")
        doc_id = str(body.get("doc_id") or "")
        if operation not in {"plan", "render"}:
            raise ValueError("operation must be 'plan' or 'render'")
        if not document_path.lower().endswith(".pptx"):
            raise ValueError("pptx_processor accepts .pptx files only")

        specialist = os.environ.get("SPECIALIST_NAME", "pptx_processor")
        subtask = job_state.subtask_id(specialist, document_path)
        job_state.mark_running(
            job_id, subtask, doc_id=doc_id, specialist=specialist,
            image_id=job_state.image_identifier(document_path), session_id=session_id,
        )
        deck = _download(document_path)
        if operation == "render":
            result = _render_selected(deck, body.get("slide_numbers"), session_id)
        else:
            result = _plan(deck, document_path, session_id, job_id, doc_id, specialist)
        job_state.mark_complete(job_id, subtask, result["artifact_uri"])
        return {
            "statusCode": 200,
            "body": json.dumps({"result": json.dumps(result, indent=2), "success": True, "session_id": session_id}),
        }
    except Exception as error:
        logger.error("PPTX processing failed: %s", error, exc_info=True)
        job_state.mark_failed(job_id, subtask, str(error))
        return {"statusCode": 400, "body": json.dumps({"error_type": "PPTX_INTAKE_FAILED", "message": str(error)})}


def _download(path: str) -> bytes:
    if path.startswith("s3://"):
        bucket, key = path.removeprefix("s3://").split("/", 1)
        return bytes(boto3.client("s3").get_object(Bucket=bucket, Key=key)["Body"].read())
    return Path(path).read_bytes()


def _plan(
    deck: bytes, document_path: str, session_id: str, job_id: str, doc_id: str, specialist: str
) -> dict[str, Any]:
    with tempfile.NamedTemporaryFile(suffix=".pptx") as source:
        source.write(deck)
        source.flush()
        presentation = Presentation(source.name)
        facts = [_inspect_slide(slide, index + 1, presentation) for index, slide in enumerate(presentation.slides)]
    s3 = boto3.client("s3")
    bucket = os.environ["OUTPUT_BUCKET"]
    pages = []
    for slide in facts:
        route = _route(slide)
        native_uri = ""
        if route["native_fidelity"] != "incomplete":
            spine = _spine(slide, document_path, session_id, route)
            if _valid_spine(spine, slide.number):
                native_uri = _put(s3, bucket, f"{session_id}/native_pptx/slide_{slide.number:03d}.xml", spine, "application/xml")
                page_task = f"{specialist}:slide_{slide.number:03d}"
                job_state.mark_running(job_id, page_task, doc_id=doc_id, specialist=specialist,
                                       image_id=job_state.image_identifier(document_path), session_id=session_id)
                job_state.mark_complete(job_id, page_task, native_uri)
            else:
                route.update({"route": "textract_ocr", "native_fidelity": "incomplete"})
                route["reasons"].append("native_spine_schema_validation_failed")
        pages.append({
            "page_number": slide.number, "slide_number": slide.number, **route,
            "content_tree_schema_valid": bool(native_uri), "native_output_uri": native_uri,
        })
    plan = {
        "source_format": "pptx", "source_document_path": document_path, "session_id": session_id,
        "page_count": len(pages), "slides": pages,
        "routing_contract": {
            "native_pptx": "native_output_uri is canonical; do not render or call visual tools",
            "textract_ocr": "render only this slide; accept OCR only when its quality gate passes",
            "hybrid": "retain native text spine and add targeted visual enrichment",
            "vision_ai": "render only this slide and apply visual analysis",
        },
    }
    uri = _put(s3, bucket, f"{session_id}/native_pptx/slide_plan.json", json.dumps(plan, indent=2), "application/json")
    return {"artifact_uri": uri, **plan}


def _inspect_slide(slide, number: int, presentation: Presentation) -> SlideFacts:
    facts = SlideFacts(number=number, slide_area=max(presentation.slide_width * presentation.slide_height, 1))
    for z_order, shape in enumerate(slide.shapes):
        _inspect_shape(shape, facts, z_order, parent_group=None)
    try:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        facts.notes = notes
    except Exception:
        facts.warnings.append("speaker_notes_unavailable")
    return facts


def _inspect_shape(shape, facts: SlideFacts, z_order: int, parent_group: str | None) -> None:
    shape_type = shape.shape_type
    geometry = {
        "x": int(shape.left), "y": int(shape.top), "width": int(shape.width), "height": int(shape.height),
        "z_order": z_order,
    }
    if shape_type == MSO_SHAPE_TYPE.GROUP:
        facts.group_count += 1
        for child_index, child in enumerate(shape.shapes):
            _inspect_shape(child, facts, z_order * 1000 + child_index, str(shape.shape_id))
        return
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            text = paragraph.text.strip()
            if text:
                facts.text_items.append({
                    "text": text, "level": int(paragraph.level), "shape_id": str(shape.shape_id),
                    "shape_name": shape.name, "parent_group": parent_group, **geometry,
                })
    if getattr(shape, "has_table", False):
        facts.table_count += 1
        for row in shape.table.rows:
            facts.text_items.append({
                "text": " | ".join(cell.text.replace("\n", " ").strip() for cell in row.cells),
                "level": 0, "shape_id": str(shape.shape_id), "shape_name": shape.name,
                "role": "table_row", "parent_group": parent_group, **geometry,
            })
    if getattr(shape, "has_chart", False):
        facts.chart_count += 1
    if shape_type == MSO_SHAPE_TYPE.PICTURE:
        facts.image_count += 1
        facts.image_area += int(shape.width) * int(shape.height)
    if shape_type == MSO_SHAPE_TYPE.LINE:
        facts.connector_count += 1
    xml = shape.element.xml
    if "drawingml/2006/diagram" in xml or "<dgm:" in xml:
        facts.smartart_count += 1
    if "oleObj" in xml or "videoFile" in xml or "audioFile" in xml:
        facts.warnings.append("embedded_object_or_media")


def _route(facts: SlideFacts) -> dict[str, Any]:
    text = facts.text
    text_chars = len(re.sub(r"\s+", "", text))
    image_ratio = min(1.0, facts.image_area / facts.slide_area)
    semantic = bool(_SEMANTIC.search(text)) or bool(
        facts.chart_count or facts.connector_count or facts.smartart_count
    )
    risk = 0
    reasons: list[str] = []
    if facts.table_count:
        risk += 8 if facts.table_count == 1 else 18
        reasons.append("native_table")
    if facts.image_count:
        risk += 8 if image_ratio < .15 else 20 if image_ratio < .5 else 35
        reasons.append("embedded_image")
    if semantic:
        risk += 35
        reasons.append("semantic_visual_shape_or_reference")
    if facts.group_count:
        risk += 10
        reasons.append("grouped_shapes")
    if facts.warnings:
        risk += 15
        reasons.extend(sorted(set(facts.warnings)))
    risk = min(100, risk)
    if text_chars < 24:
        if semantic:
            route, fidelity = "vision_ai", "incomplete"
        else:
            route, fidelity = "textract_ocr", "incomplete"
    elif semantic or risk >= 45:
        route, fidelity = "hybrid", "uncertain"
    else:
        route = "native_pptx"
        fidelity = "complete" if risk < 20 else "uncertain"
    return {
        "route": route, "native_fidelity": fidelity,
        "deterministic_review_required": route == "native_pptx" and fidelity != "complete",
        "visual_semantic_signal": semantic, "risk_score": risk,
        "reasons": reasons or ["linear_native_slide_text"],
        "evidence": {
            "native_text_characters": text_chars, "image_count": facts.image_count,
            "image_area_ratio": round(image_ratio, 3), "table_count": facts.table_count,
            "chart_count": facts.chart_count, "connector_count": facts.connector_count,
            "group_count": facts.group_count, "smartart_count": facts.smartart_count,
            "speaker_notes_characters": len(facts.notes),
        },
    }


def _spine(facts: SlideFacts, path: str, session_id: str, route: dict[str, Any]) -> str:
    root = ET.Element("unified_document", {"schema_version": "2.0", "page": str(facts.number),
                                            "source_format": "pptx", "processing_mode": "native_pptx"})
    metadata = ET.SubElement(root, "metadata")
    ET.SubElement(metadata, "source_document").text = path
    ET.SubElement(metadata, "source_image").text = path
    ET.SubElement(metadata, "session_id").text = session_id
    ET.SubElement(metadata, "native_fidelity").text = route["native_fidelity"]
    ET.SubElement(metadata, "speaker_notes").text = facts.notes
    executed = ET.SubElement(metadata, "specialists_executed")
    ET.SubElement(executed, "specialist", {"name": "pptx_processor", "s3_uri": ""})
    tree = ET.SubElement(root, "content_tree")
    section = ET.SubElement(tree, "sect", {"id": "sect_root"})
    for index, item in enumerate(facts.text_items, start=1):
        tag = "H1" if index == 1 and item["level"] == 0 else "P"
        element = ET.SubElement(section, "element", {
            "id": f"pptx_s{facts.number:03d}_{index:03d}", "tag": tag, "order": str(index),
            "page": str(facts.number), "list_level": str(item["level"]),
            "bbox": ",".join(str(item[key]) for key in ("x", "y", "width", "height")),
            "z_order": str(item["z_order"]), "provenance": "native_pptx",
        })
        ET.SubElement(element, "text").text = item["text"]
    if facts.image_count:
        figure = ET.SubElement(section, "element", {"id": f"pptx_s{facts.number:03d}_figure",
            "tag": "Figure", "order": str(len(facts.text_items) + 1), "page": str(facts.number),
            "provenance": "native_pptx_image"})
        ET.SubElement(figure, "alt_text").text = f"{facts.image_count} native picture object(s); visual meaning not inferred."
    ET.SubElement(root, "cross_reference_index")
    outputs = ET.SubElement(root, "raw_outputs")
    ET.SubElement(outputs, "output", {"specialist": "pptx_processor", "s3_uri": ""})
    ET.SubElement(root, "summary").text = f"Native PPTX extraction for slide {facts.number}; route {route['route']}."
    ET.indent(root, space="  ")
    return ET.tostring(root, encoding="unicode", xml_declaration=True)


def _render_selected(deck: bytes, slide_numbers: Any, session_id: str) -> dict[str, Any]:
    if not isinstance(slide_numbers, list) or not slide_numbers or any(
        not isinstance(number, int) or number < 1 for number in slide_numbers
    ):
        raise ValueError("render requires a non-empty slide_numbers array of positive integers")
    selected = sorted(set(slide_numbers))
    pages = []
    s3 = boto3.client("s3")
    bucket = os.environ["OUTPUT_BUCKET"]
    with tempfile.TemporaryDirectory(prefix="pptx_") as temp:
        source = Path(temp) / "source.pptx"
        source.write_bytes(deck)
        process = subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", temp, str(source)],
            capture_output=True, text=True, check=False, timeout=120,
        )
        rendered = Path(temp) / "source.pdf"
        if process.returncode or not rendered.exists():
            raise RuntimeError(f"PPTX render failed: {process.stderr[-500:]}")
        pdf = fitz.open(rendered)
        try:
            if max(selected) > len(pdf):
                raise ValueError("Requested slide number does not exist")
            for number in selected:
                pix = pdf[number - 1].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                image = Image.open(__import__("io").BytesIO(pix.tobytes("png"))).convert("RGB")
                buffer = __import__("io").BytesIO()
                image.save(buffer, format="JPEG", quality=85, optimize=True)
                key = f"temp/{session_id}/slide_{number:03d}.b64"
                s3.put_object(Bucket=bucket, Key=key, Body=base64.b64encode(buffer.getvalue()).decode(),
                              ContentType="text/plain")
                pages.append({"slide_number": number, "page_number": number, "image_path": f"s3://{bucket}/{key}"})
        finally:
            pdf.close()
    key = f"temp/{session_id}/pptx_render_manifest.json"
    result = {"rendered_slide_count": len(pages), "pages": pages}
    artifact = _put(s3, bucket, key, json.dumps(result), "application/json")
    return {"artifact_uri": artifact, **result}


def _valid_spine(xml: str, slide_number: int) -> bool:
    try:
        root = ET.fromstring(xml)
    except ET.ParseError:
        return False
    return (
        root.tag == "unified_document"
        and root.get("schema_version") == "2.0"
        and root.get("page") == str(slide_number)
        and root.find("./metadata/source_document") is not None
        and root.find("./content_tree/sect[@id='sect_root']") is not None
    )


def _put(s3: Any, bucket: str, key: str, body: str, content_type: str) -> str:
    s3.put_object(Bucket=bucket, Key=key, Body=body.encode(), ContentType=content_type)
    return f"s3://{bucket}/{key}"
